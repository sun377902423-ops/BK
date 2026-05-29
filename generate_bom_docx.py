#!/usr/bin/env python3
"""Convert 物料清单_硬件配置.md to a clean Word document on Desktop."""

import os
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DESKTOP = os.path.expanduser("~/Desktop")
MD_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "docs", "物料清单_硬件配置.md")
OUT_PATH = os.path.join(DESKTOP, "BKSYS-MED_硬件配置清单.docx")


def set_cell_shading(cell, color):
    shading = OxmlElement('w:shd')
    shading.set(qn('w:fill'), color)
    shading.set(qn('w:val'), 'clear')
    cell._tc.get_or_add_tcPr().append(shading)


def add_table_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement('w:tblPr')
    borders_elem = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        elem = OxmlElement(f'w:{edge}')
        elem.set(qn('w:val'), 'single')
        elem.set(qn('w:sz'), '4')
        elem.set(qn('w:space'), '0')
        elem.set(qn('w:color'), 'BBBBBB')
        borders_elem.append(elem)
    existing = tblPr.find(qn('w:tblBorders'))
    if existing is not None:
        tblPr.remove(existing)
    tblPr.append(borders_elem)


def main():
    with open(MD_PATH, 'r', encoding='utf-8') as f:
        md = f.read()

    doc = Document()

    style = doc.styles['Normal']
    style.font.name = '微软雅黑'
    style.font.size = Pt(10.5)
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), '微软雅黑')
    style.paragraph_format.space_after = Pt(4)
    style.paragraph_format.line_spacing = 1.25

    title = doc.add_heading('BKSYS-MED 远程医疗会诊系统', level=0)
    subtitle = doc.add_paragraph('硬件配置清单')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in subtitle.runs:
        run.font.size = Pt(14)
        run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    lines = md.split('\n')
    i = 0
    in_code_block = False
    code_buffer = []
    in_table = False
    table_rows = []

    def flush_table():
        nonlocal table_rows, in_table
        if not table_rows or len(table_rows) < 2:
            table_rows = []
            in_table = False
            return
        rows = len(table_rows)
        cols = max(len(r) for r in table_rows)
        table = doc.add_table(rows=rows, cols=cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for ri, row_data in enumerate(table_rows):
            for ci, cell_text in enumerate(row_data):
                if ci < cols:
                    cell = table.rows[ri].cells[ci]
                    cell.text = ''
                    p = cell.paragraphs[0]
                    run = p.add_run(cell_text)
                    run.font.size = Pt(9)
                    if ri == 0:
                        run.bold = True
                        set_cell_shading(cell, 'E8EAF6')
                    p.paragraph_format.space_before = Pt(1)
                    p.paragraph_format.space_after = Pt(1)
        add_table_borders(table)
        doc.add_paragraph()
        table_rows = []
        in_table = False

    while i < len(lines):
        line = lines[i].rstrip()

        if line.startswith('```'):
            if in_code_block:
                text = '\n'.join(code_buffer)
                if text.strip():
                    p = doc.add_paragraph()
                    run = p.add_run(text)
                    run.font.size = Pt(8)
                    run.font.name = 'Courier New'
                    p.paragraph_format.left_indent = Cm(0.5)
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)
                code_buffer = []
                in_code_block = False
            else:
                if in_table:
                    flush_table()
                in_code_block = True
            i += 1
            continue
        if in_code_block:
            code_buffer.append(line)
            i += 1
            continue

        if line.startswith('|') and line.endswith('|'):
            cells = [c.strip() for c in line.split('|')[1:-1]]
            if re.match(r'^\|[\s\-:]+\|', line):
                i += 1
                continue
            if in_table:
                table_rows.append(cells)
            else:
                table_rows = [cells]
                in_table = True
            i += 1
            continue
        else:
            if in_table:
                flush_table()

        if not line.strip():
            i += 1
            continue

        heading_match = re.match(r'^(#{1,4})\s+(.+)$', line)
        if heading_match:
            level = len(heading_match.group(1))
            text = heading_match.group(2)
            doc.add_heading(text, level=level)
            i += 1
            continue

        list_match = re.match(r'^(\s*)[-*+]\s+(.+)$', line)
        if list_match:
            text_val = list_match.group(2)
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.left_indent = Cm(1.0)
            p.clear()
            run = p.add_run(text_val.replace('**', ''))
            i += 1
            continue

        text = line
        p = doc.add_paragraph()
        parts = re.split(r'(\*\*[^*]+\*\*)', text)
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                run = p.add_run(part[2:-2])
                run.bold = True
            else:
                run = p.add_run(part)
        i += 1

    if in_table:
        flush_table()

    section = doc.sections[0]
    header = section.header
    hp = header.paragraphs[0]
    hp.text = 'BKSYS-MED 远程医疗会诊系统 — 硬件配置清单'
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in hp.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    doc.save(OUT_PATH)
    print(f'✅ 生成成功: {OUT_PATH}')


if __name__ == '__main__':
    main()
